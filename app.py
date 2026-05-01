 import streamlit as st
import os
import pandas as pd
from openpyxl import Workbook, load_workbook

st.title("🎓 Student Name Correction System")

# =========================
# 📌 SCHOOL LIST
# =========================
schools = [
    "YAWBRONYA JHS",
    "Prempeh",
    "Yawbronya Presby School",
    "Yaa Asantewaa Girls"
]

# =========================
# 📌 SHEETS
# =========================
sheets = [
    "KG1 SBA", "KG2 SBA",
    "P1 SBA", "P2 SBA", "P3 SBA", "P4 SBA",
    "P5 SBA", "P6 SBA",
    "JHS 1 SBA", "JHS 2 SBA", "JHS 3 SBA"
]

# =========================
# 🔽 SELECT SCHOOL & SHEET
# =========================
school = st.selectbox("Select School", schools)
sheet_name = st.selectbox("Select Class Sheet", sheets)

# =========================
# 📂 FILE NAME
# =========================
school_file = school.replace(" ", "_").lower() + ".xlsx"

# =========================
# 📁 CREATE FILE IF NOT EXISTS
# =========================
if not os.path.exists(school_file):
    wb = Workbook()
    wb.remove(wb.active)

    for s in sheets:
        ws = wb.create_sheet(title=s)
        ws.append(["Name"])

    wb.save(school_file)

# =========================
# 📄 LOAD WORKBOOK & ENSURE SHEET
# =========================
wb = load_workbook(school_file)

if sheet_name not in wb.sheetnames:
    ws = wb.create_sheet(title=sheet_name)
    ws.append(["Name"])
    wb.save(school_file)

# Load dataframe
df = pd.read_excel(school_file, sheet_name=sheet_name)

if "Name" not in df.columns:
    df["Name"] = []

# =========================
# 🔽 SELECT NAME
# =========================
name_list = df["Name"].dropna().tolist()

selected_name = st.selectbox(
    "Select Student Name",
    name_list if name_list else ["No names available"]
)

# =========================
# ✏️ UPDATE NAME
# =========================
st.subheader("✏️ Update Student Name")

new_name = st.text_input("Enter Correct Name", key="update_name")

if st.button("Update Name"):
    if selected_name == "No names available":
        st.warning("No names to update")

    elif new_name.strip() == "":
        st.warning("Enter a valid name")

    elif new_name in name_list:
        st.warning("Name already exists")

    else:
        ws = wb[sheet_name]
        updated = False

        for row in ws.iter_rows(min_row=2):
            if row[0].value == selected_name:
                row[0].value = new_name
                updated = True
                break

        wb.save(school_file)

        if updated:
            st.success(f"✅ Updated '{selected_name}' to '{new_name}'")
        else:
            st.error("Name not found")

# =========================
# ➕ ADD NEW STUDENT
# =========================
st.subheader("➕ Add New Student")

new_student = st.text_input("Enter New Student Name", key="add_name")

if st.button("Add Student"):
    if new_student.strip() == "":
        st.warning("Enter a valid name")

    elif new_student in name_list:
        st.warning("Student already exists")

    else:
        ws = wb[sheet_name]
        ws.append([new_student])
        wb.save(school_file)

        st.success(f"✅ {new_student} added successfully")

# =========================
# 🗑️ DELETE STUDENT
# =========================
st.subheader("🗑️ Delete Student")

if st.button("Delete Selected Student"):
    if selected_name == "No names available":
        st.warning("No student to delete")

    else:
        ws = wb[sheet_name]
        deleted = False

        for row in ws.iter_rows(min_row=2):
            if row[0].value == selected_name:
                ws.delete_rows(row[0].row)
                deleted = True
                break

        wb.save(school_file)

        if deleted:
            st.success(f"🗑️ '{selected_name}' deleted successfully")
        else:
            st.error("Student not found")

# =========================
# 📥 DOWNLOAD FILE
# =========================
st.subheader("📥 Download File")

with open(school_file, "rb") as file:
    st.download_button(
        "Download Updated Excel",
        data=file,
        file_name=school_file
    )
